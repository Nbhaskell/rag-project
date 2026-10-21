"""Tests for the automatic indexing (sync_docs, watch, ingest) in rag.py.

Run from the project folder with the venv active:   python test_indexing.py

Each test uses a throwaway folder and database in your temp directory, and a fake, instant
stand-in for the embedding model, so it takes seconds, does not need Ollama, and never touches
your real docs/ or chroma_db/.
"""
import hashlib
import sys
import tempfile
import threading
import time
import traceback
from pathlib import Path

import chromadb

import rag

calls = []  # how many texts each embed() call was asked to embed


def fake_embed(texts):
    calls.append(len(texts))
    return [[b / 255 for b in hashlib.sha256(t.encode()).digest()[:8]] for t in texts]


rag.embed = fake_embed


class Env:
    """A fresh docs folder and database for one test."""

    def __init__(self):
        # if Windows still holds a file open, the leftover temp folder is harmless; don't fail the test
        self.tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        root = Path(self.tmp.name)
        self.docs = root / "docs"
        self.docs.mkdir()
        rag.DOCS_DIR = self.docs
        rag.client = chromadb.PersistentClient(path=str(root / "db"))
        rag.skipped.clear()
        rag.fingerprint_cache.clear()
        rag.index_state.update(last_sync=None, last_change=None, last_summary=None)
        rag.CHUNK_SIZE = 800
        calls.clear()
        self.messages = []

    def log(self, message):
        self.messages.append(message)

    def write(self, name, text):
        path = self.docs / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding="utf-8")
        return path

    def sync(self, **kw):
        return rag.sync_docs(log=self.log, **kw)

    def chunks(self, key=None):
        got = rag.get_collection().get(include=["documents", "metadatas"])
        rows = list(zip(got["documents"], got["metadatas"]))
        return [(d, m) for d, m in rows if key is None or m["source"] == key]

    def close(self):
        try:
            rag.client._system.stop()  # let go of the database files (Windows won't delete open files)
        except Exception:
            pass
        self.tmp.cleanup()


def make(test):
    def run():
        env = Env()
        try:
            test(env)
        finally:
            env.close()
    run.__name__ = test.__name__
    return run


def check(condition, message):
    if not condition:
        raise AssertionError(message)


DOC = "# Pet Rules\n\n## Hours\nOpen 8am to 6pm.\n\n## Fees\nA bath costs $50.\n"


@make
def test_new_file_is_added(env):
    env.write("a.md", DOC)
    s = env.sync()
    check(s["added"] == ["a.md"], f"expected a.md added, got {s}")
    check(len(env.chunks("a.md")) == 2, f"expected one chunk per section (2), got {len(env.chunks('a.md'))}")
    check(all(m.get("file_hash") for _, m in env.chunks("a.md")), "chunks should carry a file_hash")


@make
def test_unchanged_files_are_not_reprocessed(env):
    env.write("a.md", DOC)
    env.sync()
    calls.clear()
    s = env.sync()
    check(s["unchanged"] == 1 and not (s["added"] or s["updated"] or s["removed"]), f"nothing should change: {s}")
    check(not calls, f"nothing should be embedded again, but embed was called {calls}")


@make
def test_changed_file_is_replaced(env):
    env.write("a.md", DOC)
    env.sync()
    env.write("a.md", DOC.replace("$50", "$75 now"))
    s = env.sync()
    check(s["updated"] == ["a.md"], f"expected an update, got {s}")
    text = " ".join(d for d, _ in env.chunks("a.md"))
    check("$75" in text and "$50" not in text, "old text must be gone and new text present")
    check(len({m["file_hash"] for _, m in env.chunks("a.md")}) == 1, "only one version should remain")


@make
def test_deleted_file_is_removed(env):
    path = env.write("a.md", DOC)
    env.write("b.md", "# Other\n\n## Topic\nSomething else.\n")
    env.sync()
    path.unlink()
    s = env.sync()
    check(s["removed"] == ["a.md"], f"expected a.md removed, got {s}")
    check(not env.chunks("a.md") and env.chunks("b.md"), "a.md chunks gone, b.md kept")


@make
def test_empty_file_is_reported_once(env):
    env.write("empty.md", "   \n")
    env.sync()
    first = [m for m in env.messages if "empty.md" in m]
    env.sync()
    env.sync()
    again = [m for m in env.messages if "empty.md" in m]
    check(len(first) == 1 and len(again) == 1, f"warning should appear once, saw {again}")
    env.write("empty.md", DOC)
    check(env.sync()["added"] == ["empty.md"], "once it has text, it should be indexed")


@make
def test_file_that_becomes_empty_is_removed(env):
    env.write("a.md", DOC)
    env.sync()
    env.write("a.md", "")
    s = env.sync()
    check(not env.chunks("a.md"), "chunks of a now-empty file should be removed")
    check(s["removed"] == ["a.md"], f"expected removal reported, got {s}")


@make
def test_unreadable_file_is_reported_once_and_keeps_old_chunks(env):
    env.write("a.docx", "this is not a real docx file")
    env.sync()
    env.sync()
    msgs = [m for m in env.messages if "a.docx" in m]
    check(len(msgs) == 1 and "could not be read" in msgs[0], f"expected one 'could not be read', saw {msgs}")


@make
def test_subfolders_and_lock_files(env):
    env.write("x/a.md", DOC)
    env.write("y/a.md", DOC)
    env.write("~$lock.docx", "temporary")
    s = env.sync()
    check(sorted(s["added"]) == ["x/a.md", "y/a.md"], f"same name in two folders must both be indexed: {s}")
    check(not env.chunks("~$lock.docx"), "Word lock files must be ignored")


@make
def test_unsettled_files_wait(env):
    env.write("a.md", DOC)
    s = env.sync(settled=set())
    check(s["waiting"] == ["a.md"] and not env.chunks(), f"unsettled file must wait: {s}")
    check(env.sync(settled={"a.md"})["added"] == ["a.md"], "settled file should be indexed")


@make
def test_changing_settings_reindexes(env):
    env.write("a.md", DOC)
    env.sync()
    rag.CHUNK_SIZE = 500
    s = env.sync()
    check(s["updated"] == ["a.md"], f"a new chunk size should re-index the file: {s}")


@make
def test_old_index_without_fingerprints_is_upgraded(env):
    env.write("a.md", DOC)
    rag.get_collection().add(ids=["a.md-0"], documents=["old chunk"], embeddings=fake_embed(["old chunk"]),
                             metadatas=[{"source": "a.md"}])
    s = env.sync()
    check(s["updated"] == ["a.md"], f"an index from before fingerprints should be upgraded: {s}")
    check(all("old chunk" != d for d, _ in env.chunks("a.md")), "the old chunk should be gone")


@make
def test_placeholder_warning(env):
    env.write("a.md", "# T\n\n## S\nCall [Name] at [phone].\n")
    s = env.sync()
    check(any("placeholder" in w for w in s["warnings"]), f"expected a placeholder warning: {s['warnings']}")


@make
def test_ingest_matches_sync(env):
    env.write("a.md", DOC)
    env.write("sub/b.md", "# B\n\n## One\nAlpha.\n\n## Two\nBeta.\n")
    env.sync()
    synced = sorted(d for d, _ in env.chunks())
    rag.ingest()
    rebuilt = sorted(d for d, _ in env.chunks())
    check(synced == rebuilt and synced, "a full rebuild and an incremental sync must give the same chunks")


@make
def test_watch_follows_edits(env):
    env.write("a.md", DOC)
    stop = threading.Event()
    thread = threading.Thread(target=rag.watch, kwargs={"stop": stop, "interval": 0.2, "log": env.log}, daemon=True)
    thread.start()

    def wait_for(condition, what, seconds=15):
        end = time.time() + seconds
        while time.time() < end:
            if condition():
                return
            time.sleep(0.1)
        raise AssertionError(f"timed out waiting for: {what}. Log: {env.messages}")

    try:
        wait_for(lambda: env.chunks("a.md"), "existing file indexed at startup")
        env.write("b.md", "# B\n\n## Topic\nBrand new file.\n")
        wait_for(lambda: env.chunks("b.md"), "new file indexed")
        env.write("a.md", DOC.replace("8am", "9am"))
        wait_for(lambda: any("9am" in d for d, _ in env.chunks("a.md")), "edited file re-indexed")
        (env.docs / "b.md").unlink()
        wait_for(lambda: not env.chunks("b.md"), "deleted file removed")
    finally:
        stop.set()
        thread.join(timeout=5)


def make_pdf(pages):
    """A tiny valid PDF with one line of text per page, so the tests can have real PDF files
    (no brackets or backslashes in the text). pypdf reads it back like any other PDF."""
    font_id = 3 + 2 * len(pages)
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(len(pages)))
    objects = ["<< /Type /Catalog /Pages 2 0 R >>", f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>"]
    for i, text in enumerate(pages):
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {4 + 2 * i} 0 R "
                       f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>")
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    objects.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out, offsets = b"%PDF-1.4\n", []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{offset:010d} 00000 n \n" for offset in offsets).encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


@make
def test_chunks_store_their_section(env):
    env.write("a.md", "# Pet Rules\n\nIntro line.\n\n## Fees\n### Baths\nA bath costs $50.\n\n## Hours\nOpen 8am.\n")
    env.sync()
    sections = {d.split("\n")[0]: m["section"] for d, m in env.chunks("a.md")}
    check(sections == {"Pet Rules": "", "Pet Rules > Fees > Baths": "Fees > Baths", "Pet Rules > Hours": "Hours"},
          f"each chunk should store its heading path without the title: {sections}")
    check(all("page" not in m for _, m in env.chunks("a.md")), "a markdown file has no page numbers")


@make
def test_pdf_chunks_know_their_page(env):
    (env.docs / "guide.pdf").write_bytes(make_pdf(["Open 8am to 6pm on weekdays", "A bath costs 50 dollars",
                                                    "Nail trims are 15 dollars"]))
    s = env.sync()
    check(s["added"] == ["guide.pdf"], f"the PDF should be indexed: {s}")
    pages = {}
    for text, meta in env.chunks("guide.pdf"):
        pages.setdefault(meta["page"], []).append(text)
        check(meta["section"] == "", f"a PDF has no headings, so no section: {meta}")
    check(sorted(pages) == [1, 2, 3], f"every page should give a chunk with its page number: {sorted(pages)}")
    check("8am" in pages[1][0] and "bath" in pages[2][0] and "Nail" in pages[3][0], f"wrong page text: {pages}")
    check(all(len(texts) == 1 for texts in pages.values()), f"a chunk must not span pages: {pages}")


@make
def test_word_headings_become_sections(env):
    from docx import Document
    doc = Document()
    doc.add_heading("Pet Rules", 1)  # a Word title has no text under it, so it gives no chunk of its own
    doc.add_heading("Hours", 2)
    doc.add_paragraph("Open 8am to 6pm.")
    doc.add_heading("Fees", 2)
    table = doc.add_table(rows=2, cols=2)
    for row, (name, price) in zip(table.rows, [("Bath", "$50"), ("Nail trim", "$30")]):
        row.cells[0].text, row.cells[1].text = name, price
    doc.save(str(env.docs / "rules.docx"))
    env.sync()
    sections = {m["section"]: d for d, m in env.chunks("rules.docx")}
    check(sorted(sections) == ["Fees", "Hours"], f"each Word heading should be a section: {sorted(sections)}")
    check("Bath | $50" in sections["Fees"] and "8am" in sections["Hours"], f"unexpected chunks: {sections}")
    check(all("page" not in m for _, m in env.chunks("rules.docx")), "a Word file has no page numbers")


@make
def test_new_index_version_reindexes(env):
    env.write("a.md", DOC)
    env.sync()
    rag.INDEX_VERSION += 1
    try:
        s = env.sync()
    finally:
        rag.INDEX_VERSION -= 1
    check(s["updated"] == ["a.md"], f"a new INDEX_VERSION should re-index the file: {s}")


TESTS = [v for k, v in sorted(globals().items()) if k.startswith("test_") and callable(v)]

if __name__ == "__main__":
    failed = 0
    for test in TESTS:
        try:
            test()
            print(f"PASS  {test.__name__}")
        except Exception:
            failed += 1
            print(f"FAIL  {test.__name__}")
            traceback.print_exc()
    print(f"\n{len(TESTS) - failed}/{len(TESTS)} passed")
    sys.exit(1 if failed else 0)
